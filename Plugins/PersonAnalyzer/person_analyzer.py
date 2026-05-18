"""
Плагин анализа людей для Watch Golem.

Функционал:
  - Определение пола и возраста (InsightFace)
  - Описание одежды (CLIP zero-shot классификация)
  - Сверка лица по базе известных лиц (InsightFace ArcFace)

Зависимости:
  pip install insightface onnxruntime open-clip-torch Pillow python-docx
"""

import os
import sys
import io
import json
import cv2
import numpy as np
from datetime import datetime

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..")))
from logger import get_logger

_log = get_logger("PersonAnalyzer")

_BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
_FACE_DB_DIR = os.path.join(_BASE_DIR, "FaceDB")

# ─── Кэши моделей ───
_insightface_app = None
_clip_model = None
_clip_preprocess = None
_tokenize = None
_torch = None
_PILImage = None
_face_database = None
_face_db_loaded = False

# ─── Маппинг пола ───
_GENDER_MAP = {0: "Женский", 1: "Мужской"}

# ─── CLIP-метки для одежды (английские → русские) ───
_UPPER_TYPES_EN = [
    "jacket", "coat", "down jacket", "windbreaker", "hoodie",
    "sweater", "t-shirt", "shirt", "blouse", "vest",
    "suit jacket", "dress", "raincoat", "tank top",
    "uniform", "polo shirt", "cardigan",
]
_UPPER_TYPES_RU = {
    "jacket": "куртка", "coat": "пальто", "down jacket": "пуховик",
    "windbreaker": "ветровка", "hoodie": "худи", "sweater": "свитер",
    "t-shirt": "футболка", "shirt": "рубашка", "blouse": "блузка",
    "vest": "жилет", "suit jacket": "пиджак", "dress": "платье",
    "raincoat": "плащ", "tank top": "майка",
    "uniform": "форма", "polo shirt": "поло", "cardigan": "кардиган",
}

_LOWER_TYPES_EN = [
    "jeans", "trousers", "shorts", "skirt", "sweatpants",
    "leggings", "dress pants", "cargo pants",
]
_LOWER_TYPES_RU = {
    "jeans": "джинсы", "trousers": "брюки", "shorts": "шорты",
    "skirt": "юбка", "sweatpants": "спортивные штаны",
    "leggings": "легинсы", "dress pants": "классические брюки",
    "cargo pants": "карго-штаны",
}

_COLORS_EN = [
    "black", "white", "gray", "blue", "light blue", "red",
    "green", "yellow", "brown", "beige", "pink",
    "purple", "orange", "dark blue", "khaki", "camouflage",
    "dark gray", "olive",
]
_COLORS_RU = {
    "black": "чёрный", "white": "белый", "gray": "серый",
    "blue": "синий", "light blue": "голубой", "red": "красный",
    "green": "зелёный", "yellow": "жёлтый", "brown": "коричневый",
    "beige": "бежевый", "pink": "розовый", "purple": "фиолетовый",
    "orange": "оранжевый", "dark blue": "тёмно-синий",
    "khaki": "хаки", "camouflage": "камуфляжный",
    "dark gray": "тёмно-серый", "olive": "оливковый",
}

# Порог cosine similarity для сверки лиц (ArcFace)
_FACE_MATCH_THRESHOLD = 0.45


# ═══════════════════════════════════════════════
#  Ленивая загрузка моделей
# ═══════════════════════════════════════════════

def _load_insightface(progress_callback=None):
    global _insightface_app
    if _insightface_app is not None:
        return _insightface_app

    if progress_callback:
        progress_callback("Загрузка InsightFace...", 0)

    import importlib.util

    # На Windows QThread не наследует пути DLL. Добавляем папку onnxruntime явно.
    _ort_spec = importlib.util.find_spec("onnxruntime")
    if _ort_spec and _ort_spec.origin and hasattr(os, "add_dll_directory"):
        try:
            os.add_dll_directory(os.path.dirname(_ort_spec.origin))
        except Exception:
            pass

    if importlib.util.find_spec("insightface") is None:
        raise ImportError(
            "InsightFace не установлен.\n"
            "Установите: pip install insightface onnxruntime"
        )
    try:
        from insightface.app import FaceAnalysis
    except ImportError as e:
        raise ImportError(
            f"Ошибка импорта InsightFace: {e}\n"
            "Проверьте установку: pip install insightface onnxruntime"
        ) from e

    try:
        _insightface_app = FaceAnalysis(
            name="buffalo_l",
            providers=['CPUExecutionProvider'],
        )
        _insightface_app.prepare(ctx_id=0, det_size=(640, 640))
    except Exception as e:
        _insightface_app = None
        raise RuntimeError(f"Ошибка инициализации InsightFace: {e}") from e

    _log.info("InsightFace загружен (buffalo_l)")
    return _insightface_app


def _load_clip(progress_callback=None):
    global _clip_model, _clip_preprocess, _tokenize, _torch, _PILImage
    if _clip_model is not None:
        return _clip_model, _clip_preprocess, _tokenize

    if progress_callback:
        progress_callback("Загрузка CLIP...", 1)

    try:
        import torch
        import open_clip
        from PIL import Image
    except ImportError as e:
        raise ImportError(
            f"CLIP зависимости не установлены: {e}\n"
            "Установите: pip install open-clip-torch Pillow"
        ) from e

    _torch = torch
    _PILImage = Image

    model, _, preprocess = open_clip.create_model_and_transforms(
        'ViT-B-32', pretrained='laion2b_s34b_b79k'
    )
    model.eval()
    tokenizer = open_clip.get_tokenizer('ViT-B-32')

    _clip_model = model
    _clip_preprocess = preprocess
    _tokenize = tokenizer
    _log.info("CLIP загружен (ViT-B-32)")
    return _clip_model, _clip_preprocess, _tokenize


# ═══════════════════════════════════════════════
#  База лиц
# ═══════════════════════════════════════════════

def _load_face_database_impl(progress_callback=None):
    """Считывает FaceDB/, вычисляет/кэширует эмбеддинги."""
    if not os.path.isdir(_FACE_DB_DIR):
        _log.info(f"Директория базы лиц не найдена: {_FACE_DB_DIR}")
        return {}

    app = _load_insightface(progress_callback)

    cache_file = os.path.join(_FACE_DB_DIR, "_embeddings_cache.json")
    cache = {}
    if os.path.isfile(cache_file):
        try:
            with open(cache_file, "r", encoding="utf-8") as f:
                cache = json.load(f)
        except Exception:
            cache = {}

    database = {}
    cache_updated = False

    for entry in sorted(os.listdir(_FACE_DB_DIR)):
        person_dir = os.path.join(_FACE_DB_DIR, entry)
        if not os.path.isdir(person_dir) or entry.startswith("_"):
            continue

        person_name = entry.replace("_", " ")
        embeddings = []

        for img_name in sorted(os.listdir(person_dir)):
            if not img_name.lower().endswith(('.jpg', '.jpeg', '.png', '.bmp')):
                continue

            img_path = os.path.join(person_dir, img_name)
            cache_key = f"{entry}/{img_name}"
            mtime = str(os.path.getmtime(img_path))

            if cache_key in cache and cache[cache_key].get("mtime") == mtime:
                emb = np.array(cache[cache_key]["embedding"], dtype=np.float32)
                embeddings.append(emb)
                continue

            img = cv2.imread(img_path)
            if img is None:
                _log.warning(f"Не удалось прочитать: {img_path}")
                continue

            faces = app.get(img)
            if not faces:
                _log.warning(f"Лицо не найдено в базе: {img_path}")
                continue

            face = max(faces, key=lambda f: (f.bbox[2] - f.bbox[0]) * (f.bbox[3] - f.bbox[1]))
            emb = face.embedding
            embeddings.append(emb)

            cache[cache_key] = {"mtime": mtime, "embedding": emb.tolist()}
            cache_updated = True

        if embeddings:
            database[person_name] = embeddings
            _log.info(f"  База лиц: '{person_name}' — {len(embeddings)} фото")

    if cache_updated:
        try:
            with open(cache_file, "w", encoding="utf-8") as f:
                json.dump(cache, f, ensure_ascii=False)
        except Exception as e:
            _log.warning(f"Ошибка сохранения кэша эмбеддингов: {e}")

    _log.info(f"База лиц: загружено {len(database)} персон")
    return database


def _get_face_database(progress_callback=None):
    global _face_database, _face_db_loaded
    if not _face_db_loaded:
        _face_database = _load_face_database_impl(progress_callback)
        _face_db_loaded = True
    return _face_database


def reload_face_database():
    """Перезагружает базу лиц (вызвать после добавления новых фото)."""
    global _face_database, _face_db_loaded
    _face_db_loaded = False
    _face_database = None


# ═══════════════════════════════════════════════
#  Сверка лица
# ═══════════════════════════════════════════════

def _match_face(embedding, database):
    if not database:
        return {"matched": False, "name": "База лиц пуста", "confidence": 0.0}

    best_name = None
    best_sim = -1.0

    for name, db_embeddings in database.items():
        for db_emb in db_embeddings:
            sim = float(np.dot(embedding, db_emb) / (
                np.linalg.norm(embedding) * np.linalg.norm(db_emb) + 1e-8
            ))
            if sim > best_sim:
                best_sim = sim
                best_name = name

    if best_sim >= _FACE_MATCH_THRESHOLD:
        return {"matched": True, "name": best_name, "confidence": best_sim}
    return {"matched": False, "name": "Неизвестный", "confidence": best_sim}


# ═══════════════════════════════════════════════
#  Описание одежды (CLIP zero-shot)
# ═══════════════════════════════════════════════

def _clip_zero_shot(image_bgr, labels, prompt_template="a photo of a person wearing {}"):
    """CLIP zero-shot классификация кропа изображения."""
    pil_img = _PILImage.fromarray(cv2.cvtColor(image_bgr, cv2.COLOR_BGR2RGB))
    img_tensor = _clip_preprocess(pil_img).unsqueeze(0)

    prompts = [prompt_template.format(lbl) for lbl in labels]
    text_tokens = _tokenize(prompts)

    with _torch.no_grad():
        img_feat = _clip_model.encode_image(img_tensor)
        txt_feat = _clip_model.encode_text(text_tokens)
        img_feat /= img_feat.norm(dim=-1, keepdim=True)
        txt_feat /= txt_feat.norm(dim=-1, keepdim=True)
        sim = (img_feat @ txt_feat.T).squeeze(0)
        probs = sim.softmax(dim=0)

    idx = probs.argmax().item()
    return labels[idx], float(probs[idx])


def _describe_clothing(image, face_bbox, progress_callback=None):
    """Описывает одежду человека по кропам тела ниже лица."""
    _load_clip(progress_callback)

    h, w = image.shape[:2]
    x1, y1, x2, y2 = [int(c) for c in face_bbox]
    face_h = y2 - y1
    face_w = x2 - x1
    center_x = (x1 + x2) // 2

    result = {}

    # ── Верхняя часть тела (ниже лица, ~2.5× высоты лица) ──
    uy1 = min(y2, h - 1)
    uy2 = min(y2 + int(face_h * 2.5), h)
    ux1 = max(0, center_x - int(face_w * 1.5))
    ux2 = min(w, center_x + int(face_w * 1.5))

    if uy2 - uy1 > 30 and ux2 - ux1 > 30:
        upper_crop = image[uy1:uy2, ux1:ux2]
        u_type, u_type_conf = _clip_zero_shot(upper_crop, _UPPER_TYPES_EN)
        u_color, u_color_conf = _clip_zero_shot(
            upper_crop, _COLORS_EN, "a photo of {} clothing"
        )
        result["upper"] = {
            "type": _UPPER_TYPES_RU.get(u_type, u_type),
            "type_confidence": u_type_conf,
            "color": _COLORS_RU.get(u_color, u_color),
            "color_confidence": u_color_conf,
        }

    # ── Нижняя часть тела ──
    ly1 = uy2 if (uy2 - uy1 > 30) else y2
    ly2 = min(ly1 + int(face_h * 3), h)
    lx1 = max(0, center_x - int(face_w * 2))
    lx2 = min(w, center_x + int(face_w * 2))

    if ly2 - ly1 > 30 and lx2 - lx1 > 30:
        lower_crop = image[ly1:ly2, lx1:lx2]
        l_type, l_type_conf = _clip_zero_shot(lower_crop, _LOWER_TYPES_EN)
        l_color, l_color_conf = _clip_zero_shot(
            lower_crop, _COLORS_EN, "a photo of {} clothing"
        )
        result["lower"] = {
            "type": _LOWER_TYPES_RU.get(l_type, l_type),
            "type_confidence": l_type_conf,
            "color": _COLORS_RU.get(l_color, l_color),
            "color_confidence": l_color_conf,
        }

    return result


# ═══════════════════════════════════════════════
#  Основной анализ
# ═══════════════════════════════════════════════

def analyze_person(image_path, progress_callback=None):
    """Анализирует человека на изображении: пол, возраст, одежда, сверка лица."""
    _log.info(f"Анализ: {os.path.basename(image_path)}")

    if not os.path.isfile(image_path):
        return {"error": f"Файл не найден: {image_path}",
                "file": os.path.basename(image_path)}

    img = cv2.imread(image_path)
    if img is None:
        return {"error": f"Не удалось прочитать: {image_path}",
                "file": os.path.basename(image_path)}

    # 1. Детекция лица + пол + возраст
    app = _load_insightface(progress_callback)
    faces = app.get(img)

    if not faces:
        return {
            "file": os.path.basename(image_path),
            "path": os.path.abspath(image_path),
            "error": "Лицо не обнаружено на изображении",
        }

    face = max(faces, key=lambda f: (f.bbox[2] - f.bbox[0]) * (f.bbox[3] - f.bbox[1]))

    # Пол
    if progress_callback:
        progress_callback("Определение пола и возраста...", 1)

    gender_val = face.get('gender', None) if isinstance(face, dict) else getattr(face, 'gender', None)
    if gender_val is None:
        gender_val = face.get('sex', None) if isinstance(face, dict) else getattr(face, 'sex', None)
    if gender_val is None:
        gender_str = "Не определён"
    elif isinstance(gender_val, str):
        gender_str = "Мужской" if gender_val.upper() == 'M' else "Женский"
    else:
        try:
            gender_str = _GENDER_MAP.get(int(gender_val), "Не определён")
        except (TypeError, ValueError):
            gender_str = "Не определён"

    # Возраст
    age_val = face.get('age', None) if isinstance(face, dict) else getattr(face, 'age', None)
    age = int(age_val) if age_val is not None else None

    det_score = float(face.det_score) if hasattr(face, 'det_score') else 1.0

    # 2. Одежда
    if progress_callback:
        progress_callback("Анализ одежды (CLIP)...", 2)
    clothing = _describe_clothing(img, face.bbox, progress_callback)

    # 3. Сверка лица по базе
    if progress_callback:
        progress_callback("Сверка по базе лиц...", 3)

    database = _get_face_database(progress_callback)
    embedding = face.embedding
    if embedding is not None:
        face_match = _match_face(embedding, database)
    else:
        face_match = {"matched": False, "name": "Эмбеддинг недоступен", "confidence": 0.0}

    if progress_callback:
        progress_callback("Готово", 4)

    result = {
        "file": os.path.basename(image_path),
        "path": os.path.abspath(image_path),
        "gender": {"gender": gender_str, "confidence": det_score},
        "age": age,
        "clothing": clothing,
        "face_match": face_match,
    }

    _log.info(
        f"Результат: {gender_str}, ~{age} лет, "
        f"лицо: {face_match['name']} ({face_match['confidence']:.2f})"
    )
    return result


def analyze_person_folder(folder_path, progress_callback=None):
    """Анализирует все изображения людей в папке."""
    extensions = ('.jpg', '.jpeg', '.png', '.bmp', '.tiff', '.gif')
    files = sorted(
        f for f in os.listdir(folder_path)
        if f.lower().endswith(extensions)
    )
    results = []
    for fname in files:
        full_path = os.path.join(folder_path, fname)
        r = analyze_person(full_path, progress_callback)
        results.append(r)
    return results


# ═══════════════════════════════════════════════
#  Текстовый отчёт
# ═══════════════════════════════════════════════

def generate_report_person(results, report_path, elapsed_seconds=None):
    """Создаёт текстовый отчёт по анализу людей."""
    os.makedirs(os.path.dirname(report_path), exist_ok=True)
    ts = datetime.now().strftime("%d-%m-%Y %H:%M:%S")

    lines = [
        "=" * 64,
        "  ОТЧЁТ: АНАЛИЗ ЛЮДЕЙ",
        f"  Дата: {ts}",
        f"  Обработано: {len(results)}",
        "=" * 64,
        "",
    ]

    for r in results:
        if "error" in r:
            lines.append(f"Файл {r.get('file', '?')} — ОШИБКА: {r['error']}")
            lines.append("")
            continue

        g = r.get("gender", {})
        gender_str = g.get("gender", "не определён")

        age = r.get("age")
        age_str = f"~{age} лет" if age is not None else "возраст не определён"

        fm = r.get("face_match", {})
        if fm.get("matched"):
            identity_str = f"{fm['name']} ({fm['confidence']:.0%})"
        else:
            identity_str = f"{fm.get('name', 'неизвестен')} ({fm.get('confidence', 0):.0%})"

        lines.append(
            f"Объектом осмотра является файл {r.get('file', '?')} на котором обнаружен "
            f"человек — пол: {gender_str}, возраст: {age_str}, личность: {identity_str}."
        )
        lines.append("")

    if elapsed_seconds is not None:
        mins, secs = divmod(int(elapsed_seconds), 60)
        lines.append("=" * 64)
        lines.append(f"  Время обработки: {mins} мин {secs} сек")

    with open(report_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))

    _log.info(f"Текстовый отчёт сохранён: {report_path}")
    return report_path


# ═══════════════════════════════════════════════
#  Word-отчёт
# ═══════════════════════════════════════════════

def generate_report_word_person(results, report_dir, elapsed_seconds=None):
    """Создаёт Word-документ с изображениями и результатами анализа."""
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from PIL import Image

    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Arial'
    style.font.size = Pt(11)

    # Заголовок
    title = doc.add_heading("Отчёт: Анализ людей", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    ts = datetime.now().strftime("%d-%m-%Y %H:%M:%S")
    p = doc.add_paragraph(f"Дата: {ts}\nОбработано изображений: {len(results)}")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for idx, r in enumerate(results):
        if idx > 0:
            doc.add_page_break()

        doc.add_heading(f"Изображение: {r.get('file', '?')}", level=2)

        # Вставка изображения
        image_path = r.get("path", "")
        if image_path and os.path.isfile(image_path):
            try:
                doc.add_picture(image_path, width=Inches(4.5))
            except ZeroDivisionError:
                try:
                    img = Image.open(image_path)
                    buf = io.BytesIO()
                    img.save(buf, format=img.format or "JPEG", dpi=(96, 96))
                    buf.seek(0)
                    doc.add_picture(buf, width=Inches(4.5))
                except Exception as e:
                    doc.add_paragraph(f"[Ошибка вставки изображения: {e}]")
            except Exception as e:
                doc.add_paragraph(f"[Ошибка вставки изображения: {e}]")

        if "error" in r:
            doc.add_paragraph(f"Ошибка: {r['error']}")
            continue

        g = r.get("gender", {})
        gender_str = g.get("gender", "не определён")

        age = r.get("age")
        age_str = f"~{age} лет" if age is not None else "возраст не определён"

        fm = r.get("face_match", {})
        if fm.get("matched"):
            identity_str = f"{fm['name']} ({fm['confidence']:.0%})"
        else:
            identity_str = f"{fm.get('name', 'неизвестен')} ({fm.get('confidence', 0):.0%})"

        verdict = doc.add_paragraph()
        verdict.add_run(
            f"Объектом осмотра является файл {r.get('file', '?')} на котором обнаружен "
            f"человек — пол: {gender_str}, возраст: {age_str}, личность: {identity_str}."
        )

    if elapsed_seconds is not None:
        doc.add_page_break()
        doc.add_heading("Информация об отчёте", level=2)
        mins, secs = divmod(int(elapsed_seconds), 60)
        t_para = doc.add_paragraph()
        t_para.add_run("Время обработки: ").bold = True
        t_para.add_run(f"{mins} мин {secs} сек")

    os.makedirs(report_dir, exist_ok=True)
    doc_path = os.path.join(report_dir, "PERSON_REPORT.docx")
    doc.save(doc_path)
    _log.info(f"Word отчёт сохранён: {doc_path}")
    return doc_path

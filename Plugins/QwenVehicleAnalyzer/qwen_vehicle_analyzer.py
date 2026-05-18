"""
Плагин анализа транспортных средств на базе Qwen2.5-VL-3B (GGUF / llama-cpp-python).

Использует квантизованную GGUF модель (~2.6 GB) вместо оригинальной (7 GB).
Работает на CPU без зависаний благодаря llama-cpp-python.

Зависимости:
  pip install llama-cpp-python --prefer-binary --extra-index-url https://abetlen.github.io/llama-cpp-python/whl/cpu
"""

import os
import json
import re
import base64
from datetime import datetime

_GGUF_DIR = os.path.join(os.path.dirname(__file__), "..", "..", "Models", "Qwen2.5-VL-3B-GGUF")
_MODEL_FILE = os.path.join(_GGUF_DIR, "Qwen2.5-VL-3B-Instruct-Q4_K_M.gguf")
_MMPROJ_FILE = os.path.join(_GGUF_DIR, "mmproj-Qwen2.5-VL-3B-Instruct-Q8_0.gguf")


class QwenVehicleAnalyzer:
    """Анализатор транспортных средств на базе Qwen2.5-VL-3B GGUF."""

    def __init__(self, model_path: str = None):
        self._model = None
        self._model_file = model_path or _MODEL_FILE
        self._mmproj_file = _MMPROJ_FILE

    def _load_model(self):
        """Ленивая загрузка GGUF модели при первом использовании."""
        if self._model is not None:
            return

        from llama_cpp import Llama
        from llama_cpp.llama_chat_format import Qwen25VLChatHandler

        if not os.path.isfile(self._model_file):
            raise FileNotFoundError(
                f"GGUF модель не найдена: {self._model_file}\n"
                "Скачайте модель из ggml-org/Qwen2.5-VL-3B-Instruct-GGUF"
            )
        if not os.path.isfile(self._mmproj_file):
            raise FileNotFoundError(
                f"mmproj файл не найден: {self._mmproj_file}\n"
                "Скачайте mmproj из ggml-org/Qwen2.5-VL-3B-Instruct-GGUF"
            )

        print("[Qwen GGUF] Загрузка модели...", flush=True)
        chat_handler = Qwen25VLChatHandler(clip_model_path=self._mmproj_file, verbose=False)
        self._model = Llama(
            model_path=self._model_file,
            chat_handler=chat_handler,
            n_ctx=2048,
            n_threads=max(4, os.cpu_count() or 4),
            n_batch=512,
            verbose=False,
        )
        print("[Qwen GGUF] Модель загружена.", flush=True)

    def _image_to_data_uri(self, image_path: str) -> str:
        """Кодирует изображение в base64 data URI."""
        ext = os.path.splitext(image_path)[1].lower().lstrip(".")
        mime = {"jpg": "jpeg", "jpeg": "jpeg", "png": "png"}.get(ext, "jpeg")
        with open(image_path, "rb") as f:
            b64 = base64.b64encode(f.read()).decode("utf-8")
        return f"data:image/{mime};base64,{b64}"

    def _query_model(self, image_path: str, prompt: str, max_tokens: int = 80) -> str:
        """Отправляет запрос модели с изображением и возвращает текстовый ответ."""
        self._load_model()
        data_uri = self._image_to_data_uri(image_path)
        return self._query_model_uri(data_uri, prompt, max_tokens)

    def _query_model_uri(self, data_uri: str, prompt: str, max_tokens: int = 80) -> str:
        """Отправляет запрос модели с уже закодированным data URI (без повторного чтения файла)."""
        self._load_model()
        response = self._model.create_chat_completion(
            messages=[
                {
                    "role": "user",
                    "content": [
                        {"type": "image_url", "image_url": {"url": data_uri}},
                        {"type": "text", "text": prompt},
                    ],
                }
            ],
            max_tokens=max_tokens,
            temperature=0.1,
        )
        return response["choices"][0]["message"]["content"].strip()


    def detect_color(self, image_path: str) -> dict:
        """Определяет цвет транспортного средства."""
        prompt = (
            "Определи цвет транспортного средства на изображении. "
            "Ответь строго в формате JSON: {\"color\": \"название цвета\", \"confidence\": число от 0 до 100}. "
            "Только JSON, без пояснений."
        )
        try:
            response = self._query_model(image_path, prompt, max_tokens=40)
            data = self._parse_json(response)
            return {
                "color": data.get("color", "Не определён"),
                "confidence": min(max(float(data.get("confidence", 0)) / 100.0, 0.0), 1.0),
            }
        except Exception as e:
            return {"color": "Не определён", "confidence": 0.0, "error": str(e)}

    def detect_model(self, image_path: str) -> dict:
        """Определяет марку и модель транспортного средства."""
        prompt = (
            "Определи марку и модель транспортного средства на изображении. "
            "Ответь строго в формате JSON: {\"brand\": \"марка\", \"model\": \"модель\", \"confidence\": число от 0 до 100}. "
            "Только JSON, без пояснений."
        )
        try:
            response = self._query_model(image_path, prompt, max_tokens=40)
            data = self._parse_json(response)
            brand = data.get("brand", "Не определена")
            model_name = data.get("model", "")
            full_name = f"{brand} {model_name}".strip() if model_name else brand
            return {
                "model": full_name,
                "confidence": min(max(float(data.get("confidence", 0)) / 100.0, 0.0), 1.0),
            }
        except Exception as e:
            return {"model": "Не определена", "confidence": 0.0, "error": str(e)}

    def detect_plate(self, image_path: str) -> dict:
        """Определяет государственный номер транспортного средства."""
        prompt = (
            "Прочитай государственный регистрационный номер транспортного средства на изображении. "
            "Ответь строго в формате JSON: {\"plate\": \"текст номера\" или null, \"confidence\": число от 0 до 100}. "
            "Если номер не виден, plate = null. Только JSON, без пояснений."
        )
        try:
            response = self._query_model(image_path, prompt, max_tokens=40)
            data = self._parse_json(response)
            plate = data.get("plate")
            confidence = min(max(float(data.get("confidence", 0)) / 100.0, 0.0), 1.0)

            if plate:
                # Очистка номера: оставляем буквы и цифры
                plate_clean = re.sub(r'[^A-ZА-Я0-9]', '', plate.upper())
                return {
                    "plate_text": plate_clean if len(plate_clean) >= 3 else None,
                    "confidence": confidence,
                    "message": f"Номер: {plate_clean} (уверенность: {confidence:.0%})" if len(plate_clean) >= 3 else "Номерной знак не распознан",
                }
            else:
                return {
                    "plate_text": None,
                    "confidence": 0.0,
                    "message": "Номерной знак не обнаружен",
                }
        except Exception as e:
            return {
                "plate_text": None,
                "confidence": 0.0,
                "message": f"Ошибка распознавания номера: {e}",
            }

    def analyze(self, image_path: str, progress_callback=None) -> dict:
        """
        Комплексный анализ транспортного средства.

        Args:
            image_path: путь к изображению
            progress_callback: функция(step_name, step_num) -> bool

        Returns:
            dict с результатами анализа
        """
        if not os.path.isfile(image_path):
            return {"error": f"Файл не найден: {image_path}"}

        if progress_callback and not progress_callback("Загрузка Qwen2.5-VL...", 0):
            return {"error": "Отменено пользователем"}

        # Загрузка модели (при первом вызове)
        try:
            self._load_model()
        except Exception as e:
            return {"error": f"Ошибка загрузки модели: {e}"}

        if progress_callback and not progress_callback("Анализ изображения...", 1):
            return {"error": "Отменено пользователем"}

        # Кодируем изображение один раз — используем для единственного запроса
        try:
            data_uri = self._image_to_data_uri(image_path)
        except Exception as e:
            return {"error": f"Ошибка чтения изображения: {e}"}

        # Один объединённый запрос вместо трёх: ~3x быстрее, точность не снижается
        combined_prompt = (
            "Проанализируй транспортное средство на изображении. "
            "Ответь строго в формате JSON (только JSON, без пояснений): "
            '{"color": "цвет", "brand": "марка", "model": "модель", '
            '"plate": "номер или null", '
            '"color_conf": число от 0 до 100, '
            '"model_conf": число от 0 до 100, '
            '"plate_conf": число от 0 до 100}.'
        )
        try:
            response_text = self._query_model_uri(data_uri, combined_prompt, max_tokens=120)
            data = self._parse_json(response_text)

            color_result = {
                "color": data.get("color", "Не определён"),
                "confidence": min(max(float(data.get("color_conf", 0)) / 100.0, 0.0), 1.0),
            }

            brand = data.get("brand", "Не определена")
            model_name = data.get("model", "")
            full_name = f"{brand} {model_name}".strip() if model_name else brand
            model_result = {
                "model": full_name,
                "confidence": min(max(float(data.get("model_conf", 0)) / 100.0, 0.0), 1.0),
            }

            plate_raw = data.get("plate")
            plate_conf = min(max(float(data.get("plate_conf", 0)) / 100.0, 0.0), 1.0)
            if plate_raw:
                plate_clean = re.sub(r'[^A-ZА-Я0-9]', '', str(plate_raw).upper())
                plate_result = {
                    "plate_text": plate_clean if len(plate_clean) >= 3 else None,
                    "confidence": plate_conf,
                    "message": (
                        f"Номер: {plate_clean} (уверенность: {plate_conf:.0%})"
                        if len(plate_clean) >= 3 else "Номерной знак не распознан"
                    ),
                }
            else:
                plate_result = {
                    "plate_text": None,
                    "confidence": 0.0,
                    "message": "Номерной знак не обнаружен",
                }
        except Exception as e:
            color_result = {"color": "Не определён", "confidence": 0.0, "error": str(e)}
            model_result = {"model": "Не определена", "confidence": 0.0, "error": str(e)}
            plate_result = {"plate_text": None, "confidence": 0.0, "message": f"Ошибка анализа: {e}"}

        if progress_callback:
            progress_callback("Готово", 3)

        return {
            "file": os.path.basename(image_path),
            "path": image_path,
            "color": color_result,
            "model": model_result,
            "plate": plate_result,
            "timestamp": datetime.now().strftime("%d-%m-%Y %H:%M:%S"),
            "engine": "Qwen2.5-VL-3B",
        }

    @staticmethod
    def _parse_json(text: str) -> dict:
        """Извлекает JSON из ответа модели."""
        # Пробуем прямой парсинг
        text = text.strip()
        try:
            return json.loads(text)
        except json.JSONDecodeError:
            pass

        # Ищем JSON-блок в тексте
        match = re.search(r'\{[^}]+\}', text, re.DOTALL)
        if match:
            try:
                return json.loads(match.group())
            except json.JSONDecodeError:
                pass

        # Ищем в блоке ```json ... ```
        match = re.search(r'```(?:json)?\s*(\{[^}]+\})\s*```', text, re.DOTALL)
        if match:
            try:
                return json.loads(match.group(1))
            except json.JSONDecodeError:
                pass

        return {}


# ═══════════════════════════════════════════════
# Функции-обёртки для совместимости с интерфейсом
# ═══════════════════════════════════════════════

# Глобальный экземпляр анализатора (загружается лениво)
_analyzer = None


def _get_analyzer() -> QwenVehicleAnalyzer:
    global _analyzer
    if _analyzer is None:
        _analyzer = QwenVehicleAnalyzer()
    return _analyzer


def analyze_vehicle_qwen(image_path: str, progress_callback=None) -> dict:
    """Анализ одного изображения транспорта через Qwen2.5-VL."""
    return _get_analyzer().analyze(image_path, progress_callback)


def analyze_vehicle_folder_qwen(folder_path: str) -> list:
    """Анализ всех изображений транспорта в папке."""
    extensions = ('.jpg', '.jpeg', '.png', '.bmp', '.tiff', '.gif')
    results = []
    for fname in sorted(os.listdir(folder_path)):
        if fname.lower().endswith(extensions):
            full_path = os.path.join(folder_path, fname)
            result = analyze_vehicle_qwen(full_path)
            results.append(result)
    return results


def generate_report_qwen(results: list, output_path: str, elapsed_seconds: float = None) -> str:
    """Формирует текстовый отчёт по результатам анализа Qwen2.5-VL."""
    lines = []
    lines.append("=" * 60)
    lines.append("  ОТЧЁТ: АНАЛИЗ ТРАНСПОРТНЫХ СРЕДСТВ (Qwen2.5-VL)")
    lines.append(f"  Дата: {datetime.now().strftime('%d-%m-%Y %H:%M:%S')}")
    lines.append(f"  Обработано изображений: {len(results)}")
    lines.append("=" * 60)
    lines.append("")

    for i, r in enumerate(results, 1):
        if "error" in r:
            lines.append(f"[{i}] {r.get('file', '?')} — ОШИБКА: {r['error']}")
            lines.append("")
            continue

        plate = r['plate'].get('plate_text') or r['plate']['message']
        lines.append(
            f"Объектом осмотра является файл {r['file']} на котором обнаружен транспорт "
            f"{r['model']['model']}, цвет {r['color']['color']}, "
            f"государственный номер {plate}."
        )
        lines.append("")

    lines.append("=" * 60)

    if elapsed_seconds is not None:
        mins, secs = divmod(int(elapsed_seconds), 60)
        lines.append(f"\n  Время обработки: {mins} мин {secs} сек")

    report_text = "\n".join(lines)

    os.makedirs(os.path.dirname(output_path) if os.path.dirname(output_path) else ".", exist_ok=True)
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(report_text)

    return output_path


def generate_report_word_qwen(results: list, output_dir: str, elapsed_seconds: float = None) -> str:
    """Формирует Word-отчёт по результатам анализа Qwen2.5-VL."""
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches
    except ImportError as exc:
        raise ImportError(
            "Для создания Word документов требуется python-docx. "
            "Установите: pip install python-docx"
        ) from exc

    os.makedirs(output_dir, exist_ok=True)

    doc = Document()
    title = doc.add_heading("Отчёт анализа транспортных средств (Qwen2.5-VL)", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    info_para = doc.add_paragraph()
    info_para.add_run("Дата: ").bold = True
    info_para.add_run(datetime.now().strftime("%d-%m-%Y %H:%M:%S"))

    count_para = doc.add_paragraph()
    count_para.add_run("Обработано изображений: ").bold = True
    count_para.add_run(str(len(results)))

    for index, result in enumerate(results, 1):
        if index > 1:
            doc.add_page_break()

        if "error" in result:
            heading = doc.add_heading(f"[{index}] {result.get('file', '?')} - ОШИБКА", level=2)
            heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
            error_para = doc.add_paragraph()
            error_para.add_run("Ошибка: ").bold = True
            error_para.add_run(result["error"])
            continue

        heading = doc.add_heading(f"[{index}] {result.get('file', '?')}", level=2)
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

        timestamp_para = doc.add_paragraph()
        timestamp_para.add_run("Дата анализа: ").bold = True
        timestamp_para.add_run(result.get("timestamp", datetime.now().strftime("%d-%m-%Y %H:%M:%S")))

        image_path = result.get("path")
        if image_path:
            if not os.path.isabs(image_path):
                image_path = os.path.abspath(image_path)
            if os.path.isfile(image_path):
                try:
                    doc.add_picture(image_path, width=Inches(5.5))
                    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                except ZeroDivisionError:
                    try:
                        import io
                        from PIL import Image as _PilImg
                        _img = _PilImg.open(image_path)
                        _buf = io.BytesIO()
                        _img.save(_buf, format='JPEG', dpi=(96, 96))
                        _buf.seek(0)
                        doc.add_picture(_buf, width=Inches(5.5))
                        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    except Exception as exc2:
                        image_para = doc.add_paragraph()
                        image_para.add_run(f"Не удалось добавить изображение: {exc2}")
                except Exception as exc:
                    image_para = doc.add_paragraph()
                    image_para.add_run(f"Не удалось добавить изображение: {exc}")
            else:
                image_para = doc.add_paragraph()
                image_para.add_run(f"Файл не найден: {image_path}")

        plate_info = result.get("plate", {})
        plate = plate_info.get("plate_text") or plate_info.get("message", "не установлен")
        color_info = result.get("color", {})
        model_info = result.get("model", {})

        verdict_para = doc.add_paragraph()
        verdict_para.add_run(
            f"Объектом осмотра является файл {result.get('file', '?')} на котором обнаружен транспорт "
            f"{model_info.get('model', 'не определена')}, цвет {color_info.get('color', 'не определён')}, "
            f"государственный номер {plate}."
        )

    doc.add_page_break()
    summary_heading = doc.add_heading("Информация об отчёте", level=2)
    summary_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    engine_para = doc.add_paragraph()
    engine_para.add_run("Анализатор: ").bold = True
    engine_para.add_run("Qwen2.5-VL-3B")

    errors_para = doc.add_paragraph()
    errors_para.add_run("Ошибок при обработке: ").bold = True
    errors_para.add_run(str(len([result for result in results if "error" in result])))

    if elapsed_seconds is not None:
        mins, secs = divmod(int(elapsed_seconds), 60)
        elapsed_para = doc.add_paragraph()
        elapsed_para.add_run("Время обработки: ").bold = True
        elapsed_para.add_run(f"{mins} мин {secs} сек")

    doc_path = os.path.join(output_dir, "Qwen25VL_Analysis_Report.docx")
    doc.save(doc_path)
    return doc_path

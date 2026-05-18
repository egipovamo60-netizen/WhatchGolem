"""
Плагин дополнения единой базы UnknownClass из внешней папки.

Сценарий:
  1. OpenSetAnalyzer автоматически пополняет постоянную базу UnknownClassDB.
  2. Данный модуль позволяет дополнительно импортировать изображения Unknown class
     из указанной пользователем папки в ту же самую базу.

Зависимости:
    pip install open-clip-torch Pillow python-docx
"""

import os
from datetime import datetime

from logger import get_logger
from Plugins.OpenSetAnalyzer.open_set_analyzer import (
    _DEFAULT_UNKNOWN_CLASS_SIMILARITY,
    supplement_unknown_class_database_from_folder,
    consolidate_unknown_class_database,
)

_log = get_logger("UnknownClassMerger")


def generate_report_unknown_class_merge(result, report_path):
    os.makedirs(os.path.dirname(report_path), exist_ok=True)

    lines = [
        "=" * 72,
        "  ОТЧЁТ: UNKNOWNCLASS DATABASE UPDATE",
        f"  Дата: {result.get('timestamp', datetime.now().strftime('%d-%m-%Y %H:%M:%S'))}",
        f"  Исходная папка: {result.get('input_root', '-')}",
        f"  База UnknownClass: {result.get('db_root', '-')}",
        f"  Изображений импортировано: {result.get('image_count', 0)}",
        f"  Пропущено файлов: {len(result.get('skipped_files', []))}",
        f"  Всего классов в базе: {result.get('class_count', 0)}",
        f"  Всего изображений в базе: {result.get('total_image_count', 0)}",
        f"  Порог схожести: {result.get('similarity_threshold', _DEFAULT_UNKNOWN_CLASS_SIMILARITY):.2f}",
        "=" * 72,
        "",
    ]

    for cluster in result.get("clusters", []):
        lines.append(
            f"{cluster['class_name']} | size={cluster['size']} | sources={len(cluster.get('source_dirs', []))}"
        )
        for item in cluster.get("items", []):
            lines.append(
                f"  - sim={float(item.get('similarity_to_cluster', 1.0)):.4f} | {item.get('source_dir', 'database')} | {os.path.basename(item.get('source_path') or item.get('path', '-'))}"
            )
        lines.append("")

    with open(report_path, "w", encoding="utf-8") as handle:
        handle.write("\n".join(lines))
    return report_path


def generate_report_json_unknown_class_merge(result, report_path):
    import json

    os.makedirs(os.path.dirname(report_path), exist_ok=True)
    with open(report_path, "w", encoding="utf-8") as handle:
        json.dump(result, handle, ensure_ascii=False, indent=2)
    return report_path


def generate_report_word_unknown_class_merge(result, report_dir):
    os.makedirs(report_dir, exist_ok=True)

    try:
        from docx import Document
    except ImportError as e:
        raise ImportError(
            "python-docx не установлен. Установите: pip install python-docx"
        ) from e

    document = Document()
    document.add_heading("UnknownClass Database Update Report", level=0)
    document.add_paragraph(f"Дата: {result.get('timestamp', datetime.now().strftime('%d-%m-%Y %H:%M:%S'))}")
    document.add_paragraph(f"Исходная папка: {result.get('input_root', '-')}" )
    document.add_paragraph(f"База UnknownClass: {result.get('db_root', '-')}")
    document.add_paragraph(f"Импортировано изображений: {result.get('image_count', 0)}")
    document.add_paragraph(f"Всего классов в базе: {result.get('class_count', 0)}")
    document.add_paragraph(f"Всего изображений в базе: {result.get('total_image_count', 0)}")
    document.add_paragraph(
        f"Порог схожести: {result.get('similarity_threshold', _DEFAULT_UNKNOWN_CLASS_SIMILARITY):.2f}"
    )

    skipped_files = result.get("skipped_files", [])
    if skipped_files:
        document.add_heading("Пропущенные файлы", level=1)
        for skipped in skipped_files:
            document.add_paragraph(
                f"{skipped.get('path', '-')}: {skipped.get('reason', '-')}",
                style="List Bullet",
            )

    document.add_heading("Классы базы UnknownClass", level=1)
    for cluster in result.get("clusters", []):
        document.add_heading(cluster.get("class_name", "UnknownClass"), level=2)
        document.add_paragraph(f"Папка: {cluster.get('dir', '-')}")
        document.add_paragraph(f"Размер класса: {cluster.get('size', 0)}")
        document.add_paragraph(f"Источник(и): {', '.join(cluster.get('source_dirs', [])) or '-'}")

        table = document.add_table(rows=1, cols=3)
        header = table.rows[0].cells
        header[0].text = "Similarity"
        header[1].text = "Source dir"
        header[2].text = "File"
        for item in cluster.get("items", []):
            row = table.add_row().cells
            row[0].text = f"{float(item.get('similarity_to_cluster', 0.0)):.4f}"
            row[1].text = item.get("source_dir", "-")
            row[2].text = os.path.basename(item.get("source_path") or item.get("path", "-"))

    output_path = os.path.join(report_dir, "UNKNOWNCLASS_DATABASE_REPORT.docx")
    document.save(output_path)
    return output_path


def analyze_unknown_class_folder(
    folder_path,
    output_dir=None,
    similarity_threshold=_DEFAULT_UNKNOWN_CLASS_SIMILARITY,
    progress_callback=None,
):
    """Импортирует изображения Unknown class из указанной папки в единую базу UnknownClassDB."""
    _log.info(f"Импорт в базу UnknownClass из папки: {folder_path}")

    result = supplement_unknown_class_database_from_folder(
        folder_path=folder_path,
        similarity_threshold=similarity_threshold,
        db_root=output_dir,
        progress_callback=progress_callback,
    )
    if result.get("error"):
        return result

    word_report_path = None
    try:
        word_report_path = generate_report_word_unknown_class_merge(
            result,
            result.get("db_root") or result.get("output_dir"),
        )
    except Exception as exc:
        _log.warning(f"Не удалось создать Word-отчёт UnknownClassMerger: {exc}")

    result["output_dir"] = result.get("db_root")
    result["merged_classes_dir"] = result.get("classes_dir")
    result["word_report_path"] = word_report_path
    return result


def merge_unknown_class_folders(
    folder_path,
    output_dir=None,
    similarity_threshold=_DEFAULT_UNKNOWN_CLASS_SIMILARITY,
    progress_callback=None,
):
    return analyze_unknown_class_folder(
        folder_path=folder_path,
        output_dir=output_dir,
        similarity_threshold=similarity_threshold,
        progress_callback=progress_callback,
    )
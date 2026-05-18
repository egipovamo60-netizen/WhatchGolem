"""Диагностика проблем с вставкой изображений в Word документ."""
import os
import sys

def diagnose_image(image_path):
    """Проводит диагностику изображения."""
    print(f"\n{'='*60}")
    print(f"Диагностика: {image_path}")
    print(f"{'='*60}\n")
    
    # 1. Проверяем, существует ли файл
    if not os.path.exists(image_path):
        print(f"❌ Файл не существует: {image_path}")
        return False
    
    print(f"✓ Файл существует")
    
    # 2. Проверяем размер
    file_size = os.path.getsize(image_path)
    print(f"✓ Размер файла: {file_size} байт ({file_size / (1024*1024):.2f} МБ)")
    
    # 3. Проверяем абсолютный путь
    abs_path = os.path.abspath(image_path)
    print(f"✓ Абсолютный путь: {abs_path}")
    print(f"  Длина пути: {len(abs_path)} символов")
    
    # 4. Пытаемся открыть с PIL
    try:
        from PIL import Image
        
        print("\n📸 Проверка PIL (Pillow):")
        img = Image.open(image_path)
        print(f"  ✓ Формат: {img.format}")
        print(f"  ✓ Размер: {img.size} пиксели")
        print(f"  ✓ Режим (mode): {img.mode}")
        
        # Пробуем verify
        try:
            img2 = Image.open(image_path)
            img2.verify()
            print(f"  ✓ Целостность проверена (verify())")
        except Exception as e:
            print(f"  ⚠ Verify не прошел: {e}")
            
    except Exception as e:
        print(f"❌ Ошибка PIL: {e}")
        return False
    
    # 5. Проверяем python-docx
    try:
        from docx import Document
        from docx.shared import Inches
        
        print("\n📄 Проверка python-docx:")
        doc = Document()
        doc.add_heading("Тестовый документ", 0)
        doc.add_paragraph(f"Попытка добавить: {os.path.basename(image_path)}")
        
        try:
            # Первый подход - с width
            doc.add_picture(image_path, width=Inches(5.5))
            print(f"  ✓ Изображение успешно добавлено в Document (width)")
        except ZeroDivisionError as e:
            print(f"  ⚠ Ошибка с width (нет DPI метаданных): {e}")
            print(f"  ℹ Попытка с height...")
            try:
                doc.add_picture(image_path, height=Inches(3.5))
                print(f"  ✓ Изображение успешно добавлено в Document (height)")
            except ZeroDivisionError:
                # Второй подход - сохраняем с установкой DPI
                print(f"  ℹ Попытка установить DPI и пересохранить...")
                try:
                    from PIL import Image
                    import io
                    
                    img = Image.open(image_path)
                    # Устанавливаем стандартный DPI если его нет
                    if img.info.get('dpi') is None or img.info['dpi'] == (0, 0):
                        dpi = (96, 96)  # стандартный DPI
                        print(f"    Устанавливаем DPI: {dpi}")
                    else:
                        dpi = img.info['dpi']
                    
                    # Сохраняем в памяти с установленным DPI
                    img_bytes = io.BytesIO()
                    img.save(img_bytes, format='JPEG', dpi=dpi)
                    img_bytes.seek(0)
                    
                    doc.add_picture(img_bytes, width=Inches(5.5))
                    print(f"  ✓ Изображение успешно добавлено в Document (с назначенным DPI)")
                except Exception as e2:
                    print(f"  ❌ Ошибка при установке DPI: {e2}")
                    raise
            
            # Пробуем сохранить
            test_doc = f"test_doc_{os.path.basename(image_path)}.docx"
            doc.save(test_doc)
            print(f"  ✓ Документ сохранен: {test_doc}")
            
            if os.path.exists(test_doc):
                size = os.path.getsize(test_doc)
                print(f"  ✓ Размер документа: {size} байт")
                os.remove(test_doc)
                print(f"  ✓ Тестовый документ удален")
            
            return True
            
        except Exception as e:
            print(f"  ❌ Ошибка при добавлении картинки: {e}")
            import traceback
            traceback.print_exc()
            return False
            
    except Exception as e:
        print(f"❌ Ошибка python-docx: {e}")
        return False


if __name__ == "__main__":
    # Тестируем все файлы в Test foto/One
    test_dir = r"z:\Coding\Watch golem stable\Test foto\One"
    
    if len(sys.argv) > 1:
        # Если указан параметр, тестируем конкретный файл
        test_file = sys.argv[1]
        if not os.path.isabs(test_file):
            test_file = os.path.join(test_dir, test_file)
        diagnose_image(test_file)
    else:
        # Тестируем все файлы
        for fname in sorted(os.listdir(test_dir)):
            if fname.lower().endswith(('.jpg', '.jpeg', '.png', '.bmp')):
                fpath = os.path.join(test_dir, fname)
                success = diagnose_image(fpath)
                print(f"\nРезультат: {'✓ OK' if success else '❌ ОШИБКА'}\n")
